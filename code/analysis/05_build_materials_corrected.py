from __future__ import annotations

import argparse
import json
import math
import platform
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import scipy
import statsmodels
import statsmodels.api as sm
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from scipy.stats import chi2_contingency, mannwhitneyu, norm
from sklearn.experimental import enable_iterative_imputer  # noqa: F401
from sklearn.impute import IterativeImputer

ROOT = Path(r"D:\Clinical_Public_Database_Research\01_ICU_Research\UO_NEQ_FORMAL_20260918\v1.1_CORRECTED")
OLD = ROOT / "CORRECTED_MANUSCRIPT_RESULTS_v1_1"
PRIVATE = ROOT / "PRIVATE_WORK"
PKG = OLD
CODE = PKG / "14_CODE"
# Row-level linkage anchors are execution-only and must never enter the shared package.
ANCHOR = PRIVATE / "MANUSCRIPT_PREP_ANCHORS"
SEED = 20260919

DIRS = [
    "01_TABLE1", "02_HARMONIZATION/raw_spotchecks", "03_EXPANDED_MODEL", "04_PRESSOR",
    "05_MISSINGNESS", "06_RESULTS_LEDGER", "07_SUPPLEMENT", "08_FIGURE_SOURCE", "09_AUDIT", "10_CODE",
]


def boolify(s: pd.Series) -> pd.Series:
    return s.astype(str).str.strip().str.lower().map({"true": 1.0, "t": 1.0, "1": 1.0, "false": 0.0, "f": 0.0, "0": 0.0})


def init_dirs() -> None:
    for d in DIRS:
        (PKG / d).mkdir(parents=True, exist_ok=True)
    ANCHOR.mkdir(parents=True, exist_ok=True)


def prepare_anchors() -> None:
    init_dirs()
    specs = {
        "mimic": {
            "src": PRIVATE / "formal_clean_mimic.csv",
            "cols": {
                "analysis_row": "analysis_row", "stay_id": "local_stay_id", "hadm_id": "local_hadm_id",
                "subject_id": "local_subject_id", "t0": "t0", "death28": "death28", "weight_kg": "weight_kg",
                "urine_ml_6h": "urine_ml_6h", "uo_ml_kg_h": "uo_ml_kg_h", "twa_neq": "twa_neq",
                "scaled_uo_neq": "scaled_uo_neq", "age": "age", "female": "female",
                "baseline_creatinine": "baseline_creatinine", "ckd_proxy": "ckd_proxy",
                "lactate_pre_t0": "lactate_pre_t0", "map_pre_t0": "map_pre_t0",
                "vent_primary": "invasive_vent_t0", "anuria": "anuria",
            },
        },
        "eicu": {
            "src": PRIVATE / "formal_clean_eicu.csv",
            "cols": {
                "analysis_row": "analysis_row", "stay_id": "local_stay_id", "hospital_stay_id": "local_hospital_stay_id",
                "hospitalid": "hospitalid", "t0": "t0", "death28": "death28", "weight_kg": "weight_kg",
                "urine_ml_6h": "urine_ml_6h", "uo_ml_kg_h": "uo_ml_kg_h", "twa_neq": "twa_neq",
                "scaled_uo_neq": "scaled_uo_neq", "age": "age", "female": "female",
                "baseline_creatinine": "baseline_creatinine", "ckd_proxy": "ckd_proxy",
                "lactate_pre_t0": "lactate_pre_t0", "map_pre_t0": "map_pre_t0",
                "vent_primary": "invasive_vent_pre_t0_proxy", "anuria": "anuria",
            },
        },
    }
    for name, spec in specs.items():
        d = pd.read_csv(spec["src"])
        d["clean_cohort"] = boolify(d["clean_cohort"]).eq(1)
        d = d.loc[d.clean_cohort & pd.to_numeric(d.death28, errors="coerce").notna()].copy()
        out = pd.DataFrame({new: d[old] for new, old in spec["cols"].items()})
        if "hospitalid" not in out:
            out["hospitalid"] = np.nan
        out.to_csv(ANCHOR / f"{name}_anchor.csv", index=False)
    print("ANCHORS_READY")


def fmt_num(x: float, digits: int = 2) -> str:
    if pd.isna(x):
        return "NA"
    return f"{x:.{digits}f}".rstrip("0").rstrip(".")


def med_iqr(s: pd.Series) -> str:
    x = pd.to_numeric(s, errors="coerce").dropna()
    if x.empty:
        return "NA"
    q = x.quantile([0.25, 0.5, 0.75])
    return f"{fmt_num(q.loc[0.5])} [{fmt_num(q.loc[0.25])}, {fmt_num(q.loc[0.75])}]"


def n_pct(s: pd.Series, denominator: int) -> str:
    x = pd.to_numeric(s, errors="coerce")
    n = int((x == 1).sum())
    observed = int(x.notna().sum())
    return f"{n} ({100*n/observed:.1f}%)" if observed else "NA"


MAIN_ROWS = [
    ("Demographics", "Age, years", "age", "cont"),
    ("Demographics", "Female sex", "female", "bin"),
    ("Demographics", "Weight, kg", "weight_kg", "cont"),
    ("Comorbidities", "Hypertension", "hypertension", "bin"),
    ("Comorbidities", "Diabetes", "diabetes", "bin"),
    ("Comorbidities", "Documented CKD or renal insufficiency", "ckd_table", "bin"),
    ("Comorbidities", "Heart failure", "heart_failure", "bin"),
    ("Comorbidities", "Ischemic heart disease or CAD", "ischemic_heart_disease", "bin"),
    ("Pre T0 severity", "Heart rate, beats/min", "heart_rate_pre_t0", "cont"),
    ("Pre T0 severity", "Systolic blood pressure, mmHg", "sbp_pre_t0", "cont"),
    ("Pre T0 severity", "MAP, mmHg", "map_pre_t0", "cont"),
    ("Pre T0 severity", "Shock index", "shock_index_pre_t0", "cont"),
    ("Pre T0 severity", "Lactate, mmol/L", "lactate_pre_t0", "cont"),
    ("Pre T0 severity", "Last measured creatinine within 24 h before T0, mg/dL", "baseline_creatinine", "cont"),
    ("Pre T0 laboratory", "White blood cell count, K/mcL", "wbc_pre_t0", "cont"),
    ("Pre T0 laboratory", "Platelet count, K/mcL", "platelet_pre_t0", "cont"),
    ("Pre T0 laboratory", "pH", "ph_pre_t0", "cont"),
    ("Support status", "Invasive ventilation documented at or before T0", "vent_table", "bin"),
    ("Early support T0 to 6 h", "Vasopressin use", "vaso_use", "bin"),
    ("Early support T0 to 6 h", "Two or more vasopressor agents", "pressor_ge2", "bin"),
    ("Exposure components", "Urine output, mL/kg/h", "uo_ml_kg_h", "cont"),
    ("Exposure components", "Average NEE, mcg/kg/min", "twa_neq", "cont"),
    ("Exposure components", "Scaled UO/NEQ", "scaled_uo_neq", "cont"),
    ("Exposure components", "Anuria", "anuria", "bin"),
    ("Outcomes", "28 day mortality", "death28", "bin"),
    ("Outcomes", "ICU length of stay, days", "icu_los_days", "cont"),
]

SUPP_EXTRA = [
    ("Comorbidities", "Chronic pulmonary disease", "chronic_pulmonary_disease", "bin"),
    ("Comorbidities", "Chronic liver disease", "chronic_liver_disease", "bin"),
    ("Comorbidities", "Malignancy", "malignancy", "bin"),
    ("Pre T0 laboratory", "Hemoglobin, g/dL", "hemoglobin_pre_t0", "cont"),
    ("Pre T0 laboratory", "Blood urea nitrogen, mg/dL", "bun_pre_t0", "cont"),
    ("Pre T0 laboratory", "Total bilirubin, mg/dL", "bilirubin_pre_t0", "cont"),
]


def prepare_df(path: Path, db: str) -> pd.DataFrame:
    d = pd.read_csv(path)
    for c in d.columns:
        if c not in {"database"}:
            if pd.api.types.is_bool_dtype(d[c]):
                d[c] = d[c].astype(int)
                continue
            vals = set(d[c].dropna().astype(str).str.strip().str.lower().unique())
            if vals and vals.issubset({"t", "f", "true", "false", "0", "1"}):
                d[c] = boolify(d[c])
                continue
            try:
                d[c] = pd.to_numeric(d[c], errors="ignore")
            except Exception:
                pass
    d["database"] = db
    d["ckd_table"] = d["ckd_proxy"] if db == "MIMIC-IV" else d["ckd_source_reviewed"]
    d["vent_table"] = d["vent_primary"] if db == "MIMIC-IV" else d["vent_source_reviewed"]
    d["shock_index_nearest_pair"] = d["shock_index_pre_t0"]
    d["shock_index_pre_t0"] = pd.to_numeric(d["heart_rate_pre_t0"], errors="coerce") / pd.to_numeric(d["sbp_pre_t0"], errors="coerce")
    return d


def quartile_table(d: pd.DataFrame, rows=MAIN_ROWS) -> pd.DataFrame:
    d = d.copy()
    d["quartile"] = pd.qcut(d.scaled_uo_neq.rank(method="first"), 4, labels=["Q1", "Q2", "Q3", "Q4"])
    groups = [("Overall", d)] + [(q, d.loc[d.quartile == q]) for q in ["Q1", "Q2", "Q3", "Q4"]]
    out = [{"database": d.database.iloc[0], "module": "Cohort", "variable": "Participants", **{name: str(len(x)) for name, x in groups}, **{f"{name}_missing_n": 0 for name, x in groups}}]
    for module, label, col, typ in rows:
        if col == "death28":
            label = "28-day all-cause mortality" if d.database.iloc[0] == "MIMIC-IV" else "28-day in-hospital mortality"
        out.append({"database": d.database.iloc[0], "module": module, "variable": label,
                    **{name: (med_iqr(x[col]) if typ == "cont" else n_pct(x[col], len(x))) for name, x in groups},
                    **{f"{name}_missing_n": int(x[col].isna().sum()) for name, x in groups}})
    return pd.DataFrame(out)


def smd_cont(a: pd.Series, b: pd.Series) -> float:
    a = pd.to_numeric(a, errors="coerce").dropna(); b = pd.to_numeric(b, errors="coerce").dropna()
    den = math.sqrt((a.var(ddof=1) + b.var(ddof=1)) / 2) if len(a)>1 and len(b)>1 else np.nan
    return float((a.mean()-b.mean())/den) if den and np.isfinite(den) else np.nan


def smd_bin(a: pd.Series, b: pd.Series) -> float:
    aa = pd.to_numeric(a, errors="coerce").dropna(); bb = pd.to_numeric(b, errors="coerce").dropna()
    pa = (aa == 1).mean(); pb = (bb == 1).mean()
    den = math.sqrt((pa*(1-pa)+pb*(1-pb))/2)
    return float((pa-pb)/den) if den else 0.0


def survival_table(d: pd.DataFrame) -> pd.DataFrame:
    surv = d.loc[d.death28 == 0]; dead = d.loc[d.death28 == 1]
    rows = [{"database":d.database.iloc[0],"module":"Cohort","variable":"Participants",
             "Overall":str(len(d)),"Survivors":str(len(surv)),"Non-survivors":str(len(dead)),
             "Overall_missing_n":0,"Survivors_missing_n":0,"Non_survivors_missing_n":0,
             "P_descriptive":np.nan,"SMD_non_survivor_vs_survivor":np.nan}]
    for module, label, col, typ in MAIN_ROWS + SUPP_EXTRA:
        if col == "death28":
            continue
        if typ == "cont":
            a = pd.to_numeric(surv[col], errors="coerce").dropna(); b = pd.to_numeric(dead[col], errors="coerce").dropna()
            p = mannwhitneyu(a, b, alternative="two-sided").pvalue if len(a) and len(b) else np.nan
            smd = smd_cont(dead[col], surv[col])
            s0,s1,so = med_iqr(surv[col]),med_iqr(dead[col]),med_iqr(d[col])
        else:
            a0 = pd.to_numeric(surv[col],errors="coerce").dropna()
            b0 = pd.to_numeric(dead[col],errors="coerce").dropna()
            tab = np.array([[(a0==1).sum(),(a0!=1).sum()],[(b0==1).sum(),(b0!=1).sum()]])
            try:
                p = chi2_contingency(tab, correction=False).pvalue if tab.min()>=0 and tab.sum()>0 else np.nan
            except ValueError:
                p = np.nan
            smd = smd_bin(dead[col],surv[col])
            s0,s1,so = n_pct(surv[col],len(surv)),n_pct(dead[col],len(dead)),n_pct(d[col],len(d))
        rows.append({"database":d.database.iloc[0],"module":module,"variable":label,"Overall":so,
                     "Survivors":s0,"Non-survivors":s1,"Overall_missing_n":int(d[col].isna().sum()),
                     "Survivors_missing_n":int(surv[col].isna().sum()),"Non_survivors_missing_n":int(dead[col].isna().sum()),
                     "P_descriptive":p,"SMD_non_survivor_vs_survivor":smd})
    return pd.DataFrame(rows)


def rcs_terms(x: np.ndarray, knots: list[float]) -> pd.DataFrame:
    k=np.asarray(knots,float); x=np.asarray(x,float); out={"ratio_linear":x}; den=(k[-1]-k[0])**2
    for j in range(2):
        out[f"ratio_rcs{j+1}"]=(np.maximum(x-k[j],0)**3-np.maximum(x-k[-2],0)**3*(k[-1]-k[j])/(k[-1]-k[-2])+np.maximum(x-k[-1],0)**3*(k[-2]-k[j])/(k[-1]-k[-2]))/den
    return pd.DataFrame(out)


def expanded_fit(d: pd.DataFrame, db: str, spec: dict) -> tuple[dict, dict]:
    miss = 100*d.shock_index_pre_t0.isna().mean()
    if miss > 20:
        raise RuntimeError(f"{db} shock-index coverage below 80%: missing={miss:.2f}%")
    d=d.copy(); d["ratio_asinh"]=np.arcsinh(d.scaled_uo_neq); d["log_nee"]=np.log(d.twa_neq)
    targets=["baseline_creatinine","map_pre_t0","shock_index_pre_t0"]
    preds=["age","female","ckd_proxy","vent_primary","death28","ratio_asinh","uo_ml_kg_h","log_nee"]
    cols=preds+targets; mat=d[cols].astype(float)
    mins=[];maxs=[]
    bounds={"baseline_creatinine":(0.05,20),"map_pre_t0":(20,200),"shock_index_pre_t0":(0.1,5)}
    for c in cols:
        lo,hi=bounds.get(c,(-np.inf,np.inf));mins.append(lo);maxs.append(hi)
    imp=IterativeImputer(max_iter=10,sample_posterior=True,random_state=SEED,min_value=np.array(mins),max_value=np.array(maxs),skip_complete=True)
    imp.fit(mat); betas=[]; covs=[]; shock_effect=[]
    shock_mu=float(d.shock_index_pre_t0.mean());shock_sd=float(d.shock_index_pre_t0.std(ddof=0))
    p=spec["mimic" if db=="MIMIC-IV" else "eicu"]
    for i in range(20):
        imp.random_state_=np.random.RandomState(SEED+i+1); filled=pd.DataFrame(imp.transform(mat),columns=cols,index=d.index)
        x=d.copy()
        for c in targets:x[c]=filled[c]
        X=pd.DataFrame({"const":1.0,"age_z":(x.age-p["age"][0])/p["age"][1],"female":x.female,
            "creatinine_log_z":(np.log1p(x.baseline_creatinine)-p["log1p_creatinine"][0])/p["log1p_creatinine"][1],
            "ckd_proxy":x.ckd_proxy,"map_z":(x.map_pre_t0-p["map"][0])/p["map"][1],"vent_primary":x.vent_primary,
            "shock_index_z":(x.shock_index_pre_t0-shock_mu)/shock_sd,
            "ratio_z_mimic":(x.ratio_asinh-spec["ratio_asinh_mean"])/spec["ratio_asinh_sd"],"anuria":x.anuria},index=x.index).astype(float)
        glm=sm.GLM(x.death28.astype(int),X,family=sm.families.Binomial())
        res=glm.fit(maxiter=200,cov_type="cluster",cov_kwds={"groups":x.hospitalid}) if db=="eICU" else glm.fit(maxiter=200)
        betas.append(np.asarray(res.params));covs.append(np.asarray(res.cov_params()));shock_effect.append(res.params["shock_index_z"])
    q=np.vstack(betas);qbar=q.mean(0);w=np.mean(covs,0);b=np.cov(q,rowvar=False,ddof=1);v=w+1.05*b
    names=X.columns.tolist();idx=names.index("ratio_z_mimic");beta=qbar[idx];se=math.sqrt(v[idx,idx])
    sidx=names.index("shock_index_z");sb=qbar[sidx];sse=math.sqrt(v[sidx,sidx])
    result={"database":db,"analysis":"POST_HOC_EXPANDED_CORE_PLUS_PRE_T0_SHOCK_INDEX","n":len(d),"deaths":int(d.death28.sum()),
            "or_per_1sd_lower_asinh_ratio":math.exp(-beta),"or_ci_low":math.exp(-beta-1.96*se),"or_ci_high":math.exp(-beta+1.96*se),
            "shock_index_or_per_1sd_higher":math.exp(sb),"shock_index_ci_low":math.exp(sb-1.96*sse),"shock_index_ci_high":math.exp(sb+1.96*sse),
            "mice_m":20,"hospital_cluster_robust":db=="eICU","model_role":"post-hoc reviewer-facing robustness; does not replace primary"}
    diag={"database":db,"variable":"shock_index_pre_t0","n":len(d),"missing_n":int(d.shock_index_pre_t0.isna().sum()),"missing_percent":miss,
          "coverage_gate_ge80":miss<=20,"imputations":20,"imputed_min_bound":0.1,"imputed_max_bound":5.0,"between_imputation_beta_sd":float(np.std(q[:,idx],ddof=1))}
    return result,diag


def set_cell_shading(cell, fill: str) -> None:
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70) -> None:
    tc=cell._tc;tcPr=tc.get_or_add_tcPr();tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:tcMar=OxmlElement('w:tcMar');tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn(f'w:{m}'))
        if node is None:node=OxmlElement(f'w:{m}');tcMar.append(node)
        node.set(qn('w:w'),str(v));node.set(qn('w:type'),'dxa')


DISPLAY_LABELS = {
    "database": "Database", "module": "Module", "variable": "Variable",
    "P_descriptive": "Descriptive P", "SMD_non_survivor_vs_survivor": "SMD",
    "primary_or_per_1sd_lower_asinh_ratio": "Primary OR (1-SD lower)",
    "primary_or_ci_low": "Primary 95% CI low", "primary_or_ci_high": "Primary 95% CI high",
    "or_per_1sd_lower_asinh_ratio": "Expanded OR (1-SD lower)",
    "or_ci_low": "Expanded 95% CI low", "or_ci_high": "Expanded 95% CI high",
    "log_or_effect_change_percent": "Change in log-OR, %", "model_role": "Model role",
    "auc_apparent": "AUC", "brier_apparent": "Brier score",
    "calibration_intercept": "Calibration intercept", "calibration_slope": "Calibration slope",
    "nonlinearity_p": "Nonlinearity P", "m4_minus_m3_auc": "M4 minus M3 AUC",
    "auc_loss_below_0_02": "AUC loss <0.02", "near_zero_exclusion_or": "Near-zero exclusion OR",
    "direction_preserved": "Direction preserved", "ratio_assumption_result": "Interpretation",
    "or_per_1sd_lower_asinh_ratio": "OR (1-SD lower)",
    "p25_minus_p75_rd": "P25-P75 risk difference", "p25_over_p75_rr": "P25/P75 risk ratio",
    "rd_ci_low": "RD 95% CI low", "rd_ci_high": "RD 95% CI high",
    "rr_ci_low": "RR 95% CI low", "rr_ci_high": "RR 95% CI high",
    "hospitalid": "Hospital ID", "landmark_n": "Landmark N",
    "uo_available_percent": "UO available, %", "nee_available_percent": "NEE available, %",
    "ratio_analyzable_n": "Ratio analyzable N", "high_measurement_coverage_hospital": "High coverage",
    "any_io_percent": "Any I/O, %", "intaketotal_percent": "Intake total, %", "nettotal_percent": "Net total, %",
}


def display_label(name: str) -> str:
    return DISPLAY_LABELS.get(name, name.replace("_", " ").replace("1sd", "1-SD").replace("ci", "CI").title())


def add_no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)


def add_title(doc: Document, title: str, note: str) -> None:
    # A manually styled normal paragraph renders consistently in Word and LibreOffice.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.font.name = 'Arial'; run.font.size = Pt(14); run.font.bold = True; run.font.color.rgb = RGBColor(0, 0, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(note); p.paragraph_format.space_after = Pt(4)
    for r in p.runs: r.font.name='Arial'; r.font.size=Pt(8)


def add_dataframe_table(doc: Document, x: pd.DataFrame, compact: bool=False) -> None:
    table=doc.add_table(rows=1,cols=len(x.columns));table.autofit=True
    hdr=table.rows[0]; add_no_split(hdr)
    trPr=hdr._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');trPr.append(rep)
    for j,c in enumerate(x.columns):
        cell=hdr.cells[j];cell.text=display_label(str(c));set_cell_shading(cell,'1F4E78');cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, 30 if compact else 55, 45, 30 if compact else 55, 45)
        for run in cell.paragraphs[0].runs:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(7 if compact else 8)
        cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i,row in x.iterrows():
        tr=table.add_row(); add_no_split(tr); cells=tr.cells
        for j,val in enumerate(row):
            sval = {"True":"Yes", "False":"No", "t":"Yes", "f":"No"}.get(str(val), str(val))
            if str(x.columns[j]) in {"description", "analysis", "ratio_assumption_result"}:
                sval = sval.replace("_", " ")
            cells[j].text=sval;cells[j].vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[j], 24 if compact else 55, 40, 24 if compact else 55, 40)
            if len(sval)<25 and j>1:cells[j].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in cells[j].paragraphs[0].runs:run.font.size=Pt(6.8 if compact else 7.5)
            if len(table.rows)%2==1:set_cell_shading(cells[j],'EAF2F8')
    table.style='Table Grid'


def write_docx(df: pd.DataFrame, path: Path, title: str, note: str, display_cols: list[str] | None=None, compact: bool=False) -> None:
    x=df.copy()
    if display_cols:x=x[[c for c in display_cols if c in x.columns]]
    for c in x.columns:
        if pd.api.types.is_float_dtype(x[c]):x[c]=x[c].map(lambda v: "NA" if pd.isna(v) else (f"{v:.3g}" if abs(v)<0.001 and v!=0 else f"{v:.3f}".rstrip('0').rstrip('.')))
    x=x.fillna("NA").astype(str)
    doc=Document();sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11);sec.page_height=Inches(8.5)
    sec.top_margin=sec.bottom_margin=Inches(0.34 if compact else 0.42);sec.left_margin=sec.right_margin=Inches(0.34 if compact else 0.42)
    styles=doc.styles
    styles['Normal'].font.name='Arial';styles['Normal'].font.size=Pt(8)
    styles['Title'].font.name='Arial';styles['Title'].font.size=Pt(14);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=RGBColor(0,0,0)
    add_title(doc,title,note)
    add_dataframe_table(doc,x,compact=compact)
    doc.save(path)


def write_panel_docx(df: pd.DataFrame, path: Path, title: str, note: str) -> None:
    x=df.copy()
    for c in x.columns:
        if pd.api.types.is_float_dtype(x[c]):
            x[c]=x[c].map(lambda v: "NA" if pd.isna(v) else (f"{v:.3g}" if abs(v)<0.001 and v!=0 else f"{v:.3f}".rstrip('0').rstrip('.')))
    x=x.fillna("NA").astype(str)
    doc=Document();sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11);sec.page_height=Inches(8.5)
    sec.top_margin=sec.bottom_margin=Inches(0.38);sec.left_margin=sec.right_margin=Inches(0.32)
    doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(7)
    for idx,(db,part) in enumerate(x.groupby('database',sort=False)):
        if idx: doc.add_page_break()
        add_title(doc, title if idx == 0 else title + " (continued)", note if idx == 0 else ("Panel B: eICU." if db == "eICU" else f"Panel: {db}."))
        p=doc.add_paragraph(f"Panel {'A' if idx == 0 else 'B'}: {db}")
        p.paragraph_format.space_after=Pt(2)
        for r in p.runs: r.font.bold=True; r.font.name='Arial'; r.font.size=Pt(9)
        add_dataframe_table(doc, part.drop(columns=['database']), compact=True)
    doc.save(path)


def build() -> None:
    init_dirs()
    mimic=prepare_df(PRIVATE/"mimic_manuscript_enriched.csv","MIMIC-IV")
    eicu=prepare_df(PRIVATE/"eicu_manuscript_enriched.csv","eICU")
    table1=pd.concat([quartile_table(mimic),quartile_table(eicu)],ignore_index=True)
    table1.to_csv(PKG/"01_TABLE1/Table1_expanded.csv",index=False,encoding="utf-8-sig")
    s1=pd.concat([survival_table(mimic),survival_table(eicu)],ignore_index=True)
    s1.to_csv(PKG/"01_TABLE1/TableS1_survival.csv",index=False,encoding="utf-8-sig")

    harmon_rows=[
     ["Age","MIMIC-IV","mimiciv_derived.icustay_detail.admission_age","ICU admission","years","DIRECT","Same construct"],
     ["Age","eICU","research_vasopressor_restart_eicu_v1.e05_sepsis_cohort.age","ICU admission","years","DIRECT","Age >89 is database-masked"],
     ["Weight","MIMIC-IV","mimiciv_derived.weight_durations.weight active at T0","T0","kg","DIRECT","Same source as denominator"],
     ["Weight","eICU","validated pressor normalization weight_kg","T0 normalization source","kg","DIRECT","Same source as denominator"],
     ["CKD","MIMIC-IV","mimiciv_derived.creatinine_baseline.ckd plus ICD raw audit","pre-existing concept","binary","PARTIAL","Frozen proxy retained in model"],
     ["CKD","eICU","diagnosis before T0 plus pastHistory renal insufficiency/failure","pre-existing documented concept","binary","PARTIAL","Source-reviewed prevalence differs from old diagnosis-only proxy"],
     ["Creatinine","MIMIC-IV","mimiciv_hosp.labevents itemid 50912","last measured in -24 h to T0","mg/dL","DIRECT_FOR_LAST_PRE_T0","Corrected after seeing v1.0 results; original admission-wide/MDRD construct retained in audit"],
     ["Creatinine","eICU","last valid lab in -24 h to T0","last pre-T0","mg/dL","DIRECT_FOR_LAST_PRE_T0","Not the same historical baseline concept as MIMIC"],
     ["Invasive ventilation","MIMIC-IV","mimiciv_derived.ventilation interval overlaps T0","at T0","binary","DIRECT","InvasiveVent or Tracheostomy"],
     ["Invasive ventilation","eICU","respiratoryCare invasive airway plus treatment MV","documented at or before T0","binary","PARTIAL","Source-reviewed descriptor; frozen model retains old treatment proxy"],
     ["SOFA","MIMIC-IV","mimiciv_derived.sepsis3","derived timestamp only","score","EXCLUDED_COMMON_MODEL","Component timestamps not recoverable"],
     ["SOFA","eICU","e05_sepsis_cohort component scores","cohort construction window","score","EXCLUDED_COMMON_MODEL","Component timestamps not retained relative to T0"],
     ["Fluid input or balance","MIMIC-IV","inputevents and output events","T0 to 6 h","mL","DESCRIPTIVE_ONLY","No source-symmetric balance reconstruction in this task"],
     ["Fluid input or balance","eICU","intakeOutput cumulative totals/cell rows","T0 to 6 h","recorded native fields","DESCRIPTIVE_ONLY","Strong hospital heterogeneity; no common expanded model"],
     ["UO","Both","frozen extraction definitions","T0 to 6 h","mL/kg/h","DIRECT_WITH_DIFFERENT_RECORD_SEMANTICS","Definitions unchanged"],
     ["NEE","Both","five-agent frozen reconstruction","T0 to 6 h","mcg/kg/min","DIRECT_WITH_INTERVAL_DIFFERENCES","MIMIC interval TWA; eICU cap240"],
     ["Outcome","MIMIC-IV","death after landmark through T0+28 d","post-landmark","binary","DIRECT","28-day all-cause mortality from linked death date"],
     ["Outcome","eICU","hospital discharge status within 28 d","post-landmark","binary","PARTIAL","Must be named 28-day in-hospital mortality"],
     ["Shock index","Both","median HR divided by median source-verified SBP","-6 h to T0","ratio","DIRECT_AFTER_SOURCE_CORRECTION","eICU Systolic child field only; vitalPeriodic first at duplicate minutes"],
    ]
    s2=pd.DataFrame(harmon_rows,columns=["variable","database","source","time_semantics","unit","harmonization_status","limitation"])
    s2.to_csv(PKG/"02_HARMONIZATION/variable_harmonization_dictionary.csv",index=False,encoding="utf-8-sig")
    s2.to_csv(PKG/"07_SUPPLEMENT/TableS2.csv",index=False,encoding="utf-8-sig")

    comp=[];audit=[]
    for d in [mimic,eicu]:
        db=d.database.iloc[0]
        for agent,col in [("Norepinephrine","ne_use"),("Epinephrine","epi_use"),("Vasopressin","vaso_use"),("Phenylephrine","pheny_use"),("Dopamine","dopa_use")]:
            comp.append({"database":db,"agent":agent,"n":int(d[col].sum()),"percent":100*d[col].mean()})
        comp += [{"database":db,"agent":"Single agent","n":int((d.pressor_agent_count==1).sum()),"percent":100*(d.pressor_agent_count==1).mean()},
                 {"database":db,"agent":"Two or more agents","n":int((d.pressor_agent_count>=2).sum()),"percent":100*(d.pressor_agent_count>=2).mean()}]
        for agent,col in [("Norepinephrine","ne_contrib"),("Epinephrine","epi_contrib"),("Vasopressin","vaso_contrib"),("Phenylephrine","pheny_contrib"),("Dopamine","dopa_contrib")]:
            audit.append({"database":db,"agent":agent,"patients_with_positive_contribution":int((d[col]>0).sum()),"median_twa_contribution":pd.to_numeric(d[col],errors="coerce").median(),"sum_twa_contribution":pd.to_numeric(d[col],errors="coerce").sum()})
        audit.append({"database":db,"agent":"Component sum vs frozen NEE","patients_with_positive_contribution":len(d),"median_twa_contribution":np.median(np.abs(d.nee_component_sum-d.twa_neq)),"sum_twa_contribution":np.max(np.abs(d.nee_component_sum-d.twa_neq))})
    comp=pd.DataFrame(comp);audit=pd.DataFrame(audit)
    comp.to_csv(PKG/"04_PRESSOR/vasopressor_composition.csv",index=False,encoding="utf-8-sig")
    audit.to_csv(PKG/"04_PRESSOR/nee_reconstruction_audit.csv",index=False,encoding="utf-8-sig")
    oldcov=pd.read_csv(OLD/"03_QC/hospital_measurement_coverage_eicu.csv")
    fluidcov=pd.read_csv(PKG/"04_PRESSOR/eicu_fluid_hospital_coverage.csv")
    hcov=oldcov.merge(fluidcov,on="hospitalid",how="outer")
    hcov.to_csv(PKG/"04_PRESSOR/hospital_coverage.csv",index=False,encoding="utf-8-sig")
    s3=comp.merge(audit,on=["database","agent"],how="outer");s3.to_csv(PKG/"07_SUPPLEMENT/TableS3.csv",index=False,encoding="utf-8-sig")

    miss=[]
    candidates=sorted(set([r[2] for r in MAIN_ROWS+SUPP_EXTRA]+["ckd_proxy","vent_primary","ckd_source_reviewed","vent_source_reviewed","past_history_rows"]))
    for d in [mimic,eicu]:
        for c in candidates:
            if c in d:
                miss.append({"database":d.database.iloc[0],"variable":c,"n":len(d),"missing_n":int(d[c].isna().sum()),"missing_percent":100*d[c].isna().mean(),"nonmissing_n":int(d[c].notna().sum())})
    s4=pd.DataFrame(miss)
    s4.to_csv(PKG/"05_MISSINGNESS/missingness_all_candidate_variables.csv",index=False,encoding="utf-8-sig")
    s4.to_csv(PKG/"07_SUPPLEMENT/TableS4.csv",index=False,encoding="utf-8-sig")

    spec=json.loads((OLD/"02_DATA/frozen_exposure_spec.json").read_text(encoding="utf-8"))
    exrows=[];diags=[]
    for d in [mimic,eicu]:
        r,g=expanded_fit(d,d.database.iloc[0],spec);exrows.append(r);diags.append(g)
    expanded=pd.DataFrame(exrows);diag=pd.DataFrame(diags)
    primary=pd.read_csv(OLD/"04_PRIMARY/primary_effects.csv")
    cmp=primary[["database","or_per_1sd_lower_asinh_ratio","or_ci_low","or_ci_high"]].rename(columns={c:f"primary_{c}" for c in ["or_per_1sd_lower_asinh_ratio","or_ci_low","or_ci_high"]}).merge(expanded,on="database")
    cmp["log_or_effect_change_percent"]=(np.log(cmp.or_per_1sd_lower_asinh_ratio)-np.log(cmp.primary_or_per_1sd_lower_asinh_ratio))/np.log(cmp.primary_or_per_1sd_lower_asinh_ratio)*100
    expanded.to_csv(PKG/"03_EXPANDED_MODEL/expanded_adjustment_results.csv",index=False,encoding="utf-8-sig")
    cmp.to_csv(PKG/"03_EXPANDED_MODEL/effect_change_vs_primary.csv",index=False,encoding="utf-8-sig")
    diag.to_csv(PKG/"03_EXPANDED_MODEL/MI_diagnostics.csv",index=False,encoding="utf-8-sig")
    cmp.to_csv(PKG/"07_SUPPLEMENT/TableS5.csv",index=False,encoding="utf-8-sig")

    s6=pd.read_csv(OLD/"05_COMPONENT/component_model_comparison.csv").merge(pd.read_csv(OLD/"05_COMPONENT/calibration.csv"),on=["database","model"],how="left",suffixes=("","_calibration"))
    s6.to_csv(PKG/"07_SUPPLEMENT/TableS6.csv",index=False,encoding="utf-8-sig")
    boot=pd.read_csv(OLD/"05_COMPONENT/bootstrap_deltas.csv")
    s8=pd.read_csv(OLD/"06_SENSITIVITY/sensitivity_results.csv")
    ratio_rows=[]
    for db in ["MIMIC-IV","eICU"]:
        p=primary.loc[primary.database.eq(db)].iloc[0]
        b=boot.loc[boot.database.eq(db)&boot.contrast.eq("M4_minus_M3")&boot.metric.eq("AUC")].iloc[0]
        s2=s8.loc[s8.database.eq(db)&s8.analysis.eq("S2_EXCLUDE_NEE_LT_0_02")].iloc[0]
        ratio_rows.append({"database":db,"primary_or_per_1sd_lower":p.or_per_1sd_lower_asinh_ratio,
                           "primary_ci_low":p.or_ci_low,"primary_ci_high":p.or_ci_high,
                           "nonlinearity_p":p.nonlinearity_p,"m4_minus_m3_auc":b.estimate,
                           "auc_loss_below_0_02_point_estimate":(-b.estimate)<.02,
                           "near_zero_exclusion_or":s2.or_per_1sd_lower_asinh_ratio,
                           "direction_preserved":s2.or_per_1sd_lower_asinh_ratio>1,
                           "ratio_assumption_result":"COMPRESSION_BENCHMARK_ONLY_NOT_NONINFERIORITY"})
    s7=pd.DataFrame(ratio_rows)
    s7.to_csv(OLD/"03_QC/ratio_assumption_results.csv",index=False,encoding="utf-8-sig")
    s7.to_csv(PKG/"07_SUPPLEMENT/TableS7.csv",index=False,encoding="utf-8-sig")
    s8.to_csv(PKG/"07_SUPPLEMENT/TableS8.csv",index=False,encoding="utf-8-sig")
    s9=hcov.copy();s9.to_csv(PKG/"07_SUPPLEMENT/TableS9.csv",index=False,encoding="utf-8-sig")

    shutil.copy2(OLD/"08_FIGURES/Figure2_source_data.csv",PKG/"08_FIGURE_SOURCE/Figure2_adjusted_risk_surface_source.csv")
    shutil.copy2(OLD/"08_FIGURES/Figure3_source_data.csv",PKG/"08_FIGURE_SOURCE/Figure3_adjusted_ratio_curves_source.csv")
    shutil.copy2(OLD/"05_COMPONENT/bootstrap_deltas.csv",PKG/"08_FIGURE_SOURCE/Figure4_component_auc_differences_source.csv")
    pd.concat([pd.read_csv(OLD/"02_DATA/cohort_flow_mimic.csv").assign(database="MIMIC-IV"),pd.read_csv(OLD/"02_DATA/cohort_flow_eicu.csv").assign(database="eICU")],ignore_index=True).to_csv(PKG/"08_FIGURE_SOURCE/Figure1_cohort_flow_source.csv",index=False)
    dist=[]
    for d in [mimic,eicu]:
        for var in ["uo_ml_kg_h","twa_neq","scaled_uo_neq"]:
            for q in np.linspace(.01,.99,99):dist.append({"database":d.database.iloc[0],"variable":var,"quantile":q,"value":d[var].quantile(q),"series":"distribution"})
        for cut in [.005,.01,.02,.05,.1]:dist.append({"database":d.database.iloc[0],"variable":"twa_neq","quantile":cut,"value":int((d.twa_neq<cut).sum()),"series":"denominator_stress_count_below_cut"})
    pd.DataFrame(dist).to_csv(PKG/"08_FIGURE_SOURCE/FigureS1_distribution_denominator_stress_source.csv",index=False)
    s8[["database","analysis","n","deaths","or_per_1sd_lower_asinh_ratio","or_ci_low","or_ci_high","status"]].to_csv(PKG/"08_FIGURE_SOURCE/FigureS2_sensitivity_forest_source.csv",index=False)

    led=[]
    sources=[("primary_effects",primary,"04_PRIMARY/primary_effects.csv"),("standardized_risks",pd.read_csv(OLD/"04_PRIMARY/standardized_risks.csv"),"04_PRIMARY/standardized_risks.csv"),("expanded_model",expanded,"03_EXPANDED_MODEL/expanded_adjustment_results.csv"),("sensitivity",s8,"07_SUPPLEMENT/TableS8.csv"),("component",s6,"07_SUPPLEMENT/TableS6.csv"),("ratio_validity",s7,"07_SUPPLEMENT/TableS7.csv")]
    for category,df,src in sources:
        idcols=[c for c in ["database","analysis","estimand","model","description"] if c in df.columns]
        for i,row in df.iterrows():
            selector=";".join(f"{c}={row[c]}" for c in idcols)
            for c,v in row.items():
                if c in idcols or not isinstance(v,(int,float,np.integer,np.floating)) or pd.isna(v):continue
                led.append({"ledger_id":f"{category}_{i}_{c}","category":category,"database":row.get("database",""),"metric":c,"value":v,"unit":"as_named","source_file":src,"source_row_selector":selector})
    ledger=pd.DataFrame(led);ledger.to_csv(PKG/"06_RESULTS_LEDGER/manuscript_numbers_master.csv",index=False,encoding="utf-8-sig")

    # ponytail: CSV is the requested auditable table format; DOCX rendering is outside this correction package.

    (PKG/"09_AUDIT/session_info.txt").write_text(f"Python {sys.version}\nPlatform {platform.platform()}\npandas {pd.__version__}\nnumpy {np.__version__}\nscipy {scipy.__version__}\nstatsmodels {statsmodels.__version__}\nSeed {SEED}\n",encoding="utf-8")
    print("MATERIALS_BUILT")


if __name__ == "__main__":
    ap=argparse.ArgumentParser();ap.add_argument("--prepare-anchors",action="store_true");ap.add_argument("--build",action="store_true");args=ap.parse_args()
    if args.prepare_anchors:prepare_anchors()
    if args.build:build()
